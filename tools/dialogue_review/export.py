"""Spreadsheet export/import for dialogue review.
  python tools/dialogue_review/export.py                 -> dialogue_review.xlsx + dialogue_review.csv
  python tools/dialogue_review/export.py --import X.xlsx -> applies edited text/status/notes back
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
import store  # noqa: E402

LINE_COLS = ["file", "id", "speaker", "context", "text", "status", "notes"]
NAME_COLS = ["kind", "id", "context", "name", "status", "notes"]


def export(out_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    lines, names = store.load_lines(), store.load_names()
    wb = Workbook()
    for title, cols, rows in (("Lines", LINE_COLS, lines), ("Names", NAME_COLS, names)):
        ws = wb.active if title == "Lines" else wb.create_sheet(title)
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        widths = {"text": 60, "context": 45, "notes": 40, "name": 30, "id": 30}
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(c, 14)
        dv = DataValidation(type="list", formula1='"draft,approved,rejected"', allow_blank=False)
        ws.add_data_validation(dv)
        col = ws.cell(1, cols.index("status") + 1).column_letter
        dv.add(f"{col}2:{col}{len(rows) + 1}")
        yellow = PatternFill("solid", fgColor="FFF2CC")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if str(row[cols.index("status")].value) != "approved":
                row[cols.index("status")].fill = yellow
        ws.freeze_panes = "A2"
    xlsx = os.path.join(out_dir, "dialogue_review.xlsx")
    wb.save(xlsx)
    csv_path = os.path.join(out_dir, "dialogue_review.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, LINE_COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(lines)
    print(f"wrote {xlsx}\nwrote {csv_path}")


def import_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    changed = 0
    current = {(l["file"], l["id"]): l for l in store.load_lines()}
    ws = wb["Lines"]
    cols = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(cols, row))
        key = (r.get("file"), r.get("id"))
        if key not in current:
            continue
        cur = current[key]
        new = {k: ("" if r.get(k) is None else str(r.get(k))) for k in ("text", "status", "notes")}
        if any(new[k] != str(cur.get(k, "")) for k in new):
            store.update_line(key[0], key[1], **new)
            changed += 1
    if "Names" in wb.sheetnames:
        ws = wb["Names"]
        cols = [c.value for c in ws[1]]
        names = {(n["kind"], n["id"]): n for n in store.load_names()}
        for row in ws.iter_rows(min_row=2, values_only=True):
            r = dict(zip(cols, row))
            cur = names.get((r.get("kind"), r.get("id")))
            if cur and (str(r.get("name")) != cur["name"] or str(r.get("status")) != cur["status"]):
                store.update_name(r["kind"], r["id"], name=str(r.get("name")), status=str(r.get("status")))
                changed += 1
    print(f"applied {changed} change(s)")


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "--import":
        import_xlsx(sys.argv[2])
    else:
        export(sys.argv[1] if len(sys.argv) > 1 else os.getcwd())
